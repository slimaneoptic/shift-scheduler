"""Excel roster: colored shift grid, coverage check and per-person totals."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from model import DAYS, SHIFTS, WEEKEND, Roster, Rules

TINT = {"M": "D4E4F7", "E": "FBE1D6", "N": "D1EFE4"}  # light versions of the app's shift colors
OFF, GAP, HEAD = "F0F2F4", "F8D7D7", "1B3A57"


def _head(ws) -> None:
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=HEAD)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")


def workbook(rules: Rules, roster: Roster) -> bytes:
    names = [e.name for e in rules.staff]
    grid = pd.DataFrame(
        [[SHIFTS[roster.assign[(e, d)]][0] if (e, d) in roster.assign else ("Leave" if (e, d) in rules.off else "Off") for d in DAYS] for e in names],
        index=names, columns=DAYS,
    )
    grid["Shifts"] = [sum(1 for d in DAYS if (e, d) in roster.assign) for e in names]
    cover = pd.DataFrame(
        {d: [f"{sum(1 for e in names if roster.assign.get((e, d)) == s)}/{rules.need.get((d, s), 0)}" for s in SHIFTS] for d in DAYS},
        index=[SHIFTS[s][0] for s in SHIFTS],
    )
    people = pd.DataFrame(
        [{"name": e.name, "senior": "yes" if e.senior else "", "shifts": grid.loc[e.name, "Shifts"], "contract": f"{e.min_shifts}–{e.max_shifts}",
          "nights": sum(1 for d in DAYS if roster.assign.get((e.name, d)) == "N"),
          "weekend shifts": sum(1 for d in DAYS if d in WEEKEND and (e.name, d) in roster.assign)} for e in rules.staff]
    )
    gaps = {(d, s) for d, s, _ in roster.shortages}

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        grid.to_excel(xl, sheet_name="Roster", index_label="Employee")
        ws = xl.sheets["Roster"]
        _head(ws)
        code = {v[0]: k for k, v in SHIFTS.items()}
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=1 + len(DAYS)):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=TINT.get(code.get(cell.value), OFF))
                cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions["A"].width = 14
        for i in range(2, len(DAYS) + 3):
            ws.column_dimensions[get_column_letter(i)].width = 10
        ws.freeze_panes = "B2"

        cover.to_excel(xl, sheet_name="Coverage", index_label="Shift")
        ws = xl.sheets["Coverage"]
        _head(ws)
        for r_i, s in enumerate(SHIFTS, start=2):
            for c_i, d in enumerate(DAYS, start=2):
                cell = ws.cell(row=r_i, column=c_i)
                cell.alignment = Alignment(horizontal="center")
                if (d, s) in gaps:
                    cell.fill = PatternFill("solid", fgColor=GAP)
                    cell.font = Font(bold=True, color="9B1C1C")
        ws.column_dimensions["A"].width = 12

        people.to_excel(xl, sheet_name="Per person", index=False)
        _head(xl.sheets["Per person"])
        xl.sheets["Per person"].column_dimensions["A"].width = 14
    return buf.getvalue()
